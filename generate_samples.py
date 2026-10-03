import os
from pathlib import Path
import docx
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable

SAMPLES_DIR = Path(__file__).resolve().parent / "samples"
SAMPLES_DIR.mkdir(parents=True, exist_ok=True)


def create_strong_swe_docx():
    doc = docx.Document()

    # Header
    doc.add_heading("David Chen", level=0)
    p = doc.add_paragraph("San Francisco, CA | david.chen@email.com | +1 (555) 234-5678 | linkedin.com/in/davidchen-dev | github.com/davidchen-dev")

    # Professional Summary
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(
        "Performance-driven Senior Full Stack Engineer with over 6 years of experience architecting distributed cloud systems, "
        "asynchronous microservices, and modern web applications. Proven track record of scaling high-throughput APIs and leading engineering teams."
    )

    # Technical Skills
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Programming Languages: Python, JavaScript, TypeScript, SQL, HTML, CSS, Bash")
    doc.add_paragraph("Frameworks & Libraries: FastAPI, Django, Flask, React, Next.js, Tailwind CSS")
    doc.add_paragraph("Databases & Caching: PostgreSQL, Redis, MongoDB, SQLite")
    doc.add_paragraph("Cloud & DevOps: Docker, Kubernetes, AWS, CI/CD, GitHub Actions, Linux, Nginx")
    doc.add_paragraph("Tools & Practices: Git, REST APIs, Microservices, Agile/Scrum, Postman")

    # Professional Experience
    doc.add_heading("Professional Experience", level=1)

    doc.add_heading("Senior Backend Engineer | CloudScale Tech (2022 - Present)", level=2)
    doc.add_paragraph("• Architected asynchronous FastAPI microservices and PostgreSQL database schemas, reducing API response latency by 42% across 150,000+ daily active users.", style='List Bullet')
    doc.add_paragraph("• Engineered containerized deployment workflows using Docker, Kubernetes, and GitHub Actions CI/CD, expediting deployment cycles from 4 hours to 12 minutes.", style='List Bullet')
    doc.add_paragraph("• Implemented distributed caching layers with Redis and optimized SQL queries, scaling system throughput by 3.5x under peak seasonal traffic.", style='List Bullet')
    doc.add_paragraph("• Mentored a team of 5 junior engineers through code reviews, technical architecture sessions, and automated unit testing.", style='List Bullet')

    doc.add_heading("Full Stack Developer | Nexa Solutions (2019 - 2022)", level=2)
    doc.add_paragraph("• Built dynamic interactive customer dashboards using React, TypeScript, and Tailwind CSS, increasing user engagement by 28%.", style='List Bullet')
    doc.add_paragraph("• Developed secure RESTful APIs using Python, Django REST Framework, and PostgreSQL with robust OAuth authentication.", style='List Bullet')
    doc.add_paragraph("• Automated regression and integration test pipelines with pytest and Postman, achieving 94% test coverage and eliminating critical production defects.", style='List Bullet')

    # Key Projects
    doc.add_heading("Key Projects", level=1)
    doc.add_heading("Distributed Task Orchestrator (github.com/davidchen-dev/task-queue)", level=2)
    doc.add_paragraph("• Spearheaded an open-source distributed job queue in Python and Redis handling up to 10,000 tasks per second with automatic retries and dead-letter queues.")

    # Education
    doc.add_heading("Education", level=1)
    doc.add_paragraph("Bachelor of Science in Computer Science | University of California, Berkeley (2015 - 2019)")

    file_path = SAMPLES_DIR / "sample_strong_swe.docx"
    doc.save(str(file_path))
    print(f"Generated: {file_path}")


def create_weak_candidate_pdf():
    file_path = SAMPLES_DIR / "sample_weak_candidate.pdf"
    doc = SimpleDocTemplate(str(file_path), pagesize=letter, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()

    story = []
    
    # Title with unprofessional email and bad info
    story.append(Paragraph("<b>Johnny Developer</b>", styles["Title"]))
    story.append(Paragraph("Email: coolgamerboy99@gmail.com | Photo Attached: See photo.jpg", styles["Normal"]))
    story.append(Paragraph("Date of Birth: 14 May 1998 | Marital Status: Single | Gender: Male", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>About Me</b>", styles["Heading2"]))
    story.append(Paragraph("I am a hardworking results-oriented team player and go-getter who thinks outside the box with synergy. I am looking for a job where I can use my skills. I am a fast learner and a coding ninja.", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>Technical Skills</b>", styles["Heading2"]))
    story.append(Paragraph("Python, Kubernetes, TensorFlow, PyTorch, GraphQL, Java, Docker, AWS, React", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>Experiance</b>", styles["Heading2"]))  # Spelling error
    story.append(Paragraph("<b>Junior Helper | Web Co (2022 - 2023)</b>", styles["Normal"]))
    story.append(Paragraph("• Responsible for working on backend tasks and helped with bug fixes.", styles["Normal"]))
    story.append(Paragraph("• Assisted with the the database and worked with clients.", styles["Normal"]))  # Repeated words: the the
    story.append(Paragraph("• I handled software enginering tasks as part of the team and did mangement.", styles["Normal"]))  # Spelling errors: enginering, mangement, pronoun I
    story.append(Paragraph("• Worked on writing code in Python and attended meetings.", styles["Normal"]))  # No metrics
    story.append(Spacer(1, 10))

    story.append(Paragraph("This is a very long paragraph of continuous text where I describe my daily duties at my previous job without using any bullet points whatsoever. I was involved in many meetings and helped out wherever I could, doing whatever tasks were assigned to me by the supervisor and working late when needed, but without any specific measurable metrics or data to support my claims.", styles["Normal"]))

    doc.build(story)
    print(f"Generated: {file_path}")


def create_non_tech_pdf():
    file_path = SAMPLES_DIR / "sample_non_tech.pdf"
    doc = SimpleDocTemplate(str(file_path), pagesize=letter, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()

    story = []
    story.append(Paragraph("<b>Sarah Jenkins</b>", styles["Title"]))
    story.append(Paragraph("Chicago, IL | sarah.jenkins@gmail.com | +1 (555) 876-5432 | linkedin.com/in/sarahjenkins", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>Professional Summary</b>", styles["Heading2"]))
    story.append(Paragraph("Strategic Digital Marketing & Brand Specialist with 5 years of experience leading multi-channel campaigns, client communication, and creative content strategy.", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>Core Competencies</b>", styles["Heading2"]))
    story.append(Paragraph("Digital Marketing, Content Strategy, Communication, Leadership, Collaboration, Problem Solving, Social Media Marketing, CRM Management, Brand Strategy", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>Professional Experience</b>", styles["Heading2"]))
    story.append(Paragraph("<b>Marketing Campaign Lead | Horizon Media (2021 - Present)</b>", styles["Normal"]))
    story.append(Paragraph("• Spearheaded omni-channel marketing campaigns reaching over 500,000 potential customers, generating a 35% increase in qualified sales leads.", styles["Normal"]))
    story.append(Paragraph("• Led cross-functional collaboration across graphic design and copy teams to execute 12 product launches on schedule.", styles["Normal"]))
    story.append(Paragraph("• Analyzed user engagement metrics and optimized monthly budget allocations, reducing acquisition costs by 18%.", styles["Normal"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>Education</b>", styles["Heading2"]))
    story.append(Paragraph("Bachelor of Arts in Communications | Northwestern University (2016 - 2020)", styles["Normal"]))

    doc.build(story)
    print(f"Generated: {file_path}")


if __name__ == "__main__":
    create_strong_swe_docx()
    create_weak_candidate_pdf()
    create_non_tech_pdf()
