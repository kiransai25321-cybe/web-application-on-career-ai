import os
import re
import tempfile
from typing import Dict, Tuple, List, Optional
import pdfplumber
import docx

SECTION_PATTERNS = {
    "summary": re.compile(r"^(professional\s+summary|summary|profile|about\s+me|objective)", re.I),
    "skills": re.compile(r"^(technical\s+skills|core\s+competencies|skills|technologies|tools\s+&\s+technologies|stack)", re.I),
    "experience": re.compile(r"^(work\s+experience|professional\s+experience|experience|experiance|employment\s+history|career\s+history)", re.I),
    "projects": re.compile(r"^(projects|personal\s+projects|academic\s+projects|key\s+projects|portfolio)", re.I),
    "education": re.compile(r"^(education|academic\s+background|qualifications)", re.I),
    "certifications": re.compile(r"^(certifications|licenses|courses|awards|achievements)", re.I),
}


class ResumeParser:
    @staticmethod
    def parse_pdf(file_bytes: bytes) -> Tuple[str, Dict[str, str], Dict[str, any]]:
        """
        Parses PDF file bytes using pdfplumber.
        Returns (full_text, sections_dict, metadata_dict).
        """
        full_text_pages: List[str] = []
        has_tables = False
        page_count = 0

        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name

        try:
            with pdfplumber.open(tmp_path) as pdf:
                page_count = len(pdf.pages)
                for page in pdf.pages:
                    text = page.extract_text() or ""
                    tables = page.extract_tables()
                    if tables and len(tables) > 0:
                        has_tables = True
                    full_text_pages.append(text)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

        full_text = "\n".join(full_text_pages)
        sections = ResumeParser.split_into_sections(full_text)
        metadata = {
            "page_count": page_count,
            "has_tables": has_tables,
            "word_count": len(full_text.split()),
            "character_count": len(full_text),
        }
        return full_text, sections, metadata

    @staticmethod
    def parse_docx(file_bytes: bytes) -> Tuple[str, Dict[str, str], Dict[str, any]]:
        """
        Parses DOCX file bytes using python-docx.
        Returns (full_text, sections_dict, metadata_dict).
        """
        paragraphs: List[str] = []
        has_tables = False

        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name

        try:
            doc = docx.Document(tmp_path)
            for p in doc.paragraphs:
                if p.text.strip():
                    paragraphs.append(p.text.strip())
            
            if len(doc.tables) > 0:
                has_tables = True
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            paragraphs.append(row_text)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

        full_text = "\n".join(paragraphs)
        sections = ResumeParser.split_into_sections(full_text)
        word_count = len(full_text.split())
        metadata = {
            "page_count": max(1, (word_count // 350) + 1),
            "has_tables": has_tables,
            "word_count": word_count,
            "character_count": len(full_text),
        }
        return full_text, sections, metadata

    @staticmethod
    def split_into_sections(text: str) -> Dict[str, str]:
        """
        Splits resume text into recognized standard sections.
        """
        lines = text.splitlines()
        sections: Dict[str, List[str]] = {
            "header": [],
            "summary": [],
            "skills": [],
            "experience": [],
            "projects": [],
            "education": [],
            "certifications": [],
            "other": []
        }

        current_section = "header"

        for line in lines:
            trimmed = line.strip()
            if not trimmed:
                continue

            # Check if line is a section heading
            matched_section = None
            if len(trimmed) < 45 and not trimmed.endswith("."):
                # Strip leading/trailing symbols
                clean_heading = re.sub(r"^[^a-zA-Z0-9]+|[^a-zA-Z0-9]+$", "", trimmed)
                for sec_name, pattern in SECTION_PATTERNS.items():
                    if pattern.search(clean_heading):
                        matched_section = sec_name
                        break

            if matched_section:
                current_section = matched_section
            else:
                sections[current_section].append(trimmed)

        return {k: "\n".join(v).strip() for k, v in sections.items() if v}
